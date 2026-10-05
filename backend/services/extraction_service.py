"""
extraction_service.py
Extracts text content from PDF, DOCX, PPTX, TXT, and image files.
Returns a list of {page, text, source, ocr_used, ocr_confidence} dicts.
"""

import os
import hashlib
from typing import List, Dict

import re

# Optional OCR support
try:
    import pytesseract
    from PIL import Image as PILImage
    TESSERACT_AVAILABLE = True

    # Check common Windows paths for tesseract binary if not in PATH
    tesseract_candidates = [
        os.environ.get("TESSERACT_CMD", ""),
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        os.path.expandvars(r"%LOCALAPPDATA%\Programs\Tesseract-OCR\tesseract.exe"),
    ]
    for cand in tesseract_candidates:
        if cand and os.path.isfile(cand):
            pytesseract.pytesseract.tesseract_cmd = cand
            break
except ImportError:
    TESSERACT_AVAILABLE = False


def clean_document_text(text: str) -> str:
    """
    Clean extracted text:
    - remove null bytes and non-printable control characters (except newline/tab)
    - normalize CRLF / CR to LF
    - normalize multiple horizontal whitespaces to a single space
    - remove excessive blank lines (>2 newlines collapsed to 2)
    - strip leading/trailing whitespace
    """
    if not text:
        return ""
    # Remove null characters and control characters
    text = text.replace("\x00", "")
    text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    # Normalize line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Normalize multiple horizontal spaces/tabs on lines
    text = re.sub(r"[ \t]+", " ", text)
    # Strip spaces before newlines
    text = re.sub(r" +\n", "\n", text)
    # Normalize excessive newlines (max 2 consecutive)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _text_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def extract_pdf(file_path: str) -> List[Dict]:
    """Extract text from a PDF file, page by page.
    Falls back to OCR for image-only pages (< 30 chars of selectable text)."""
    import pymupdf as fitz

    pages = []
    doc = fitz.open(file_path)

    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text").strip()
        ocr_used = False
        ocr_confidence = 1.0

        # If page has no meaningful text, try OCR
        if len(text) < 30 and TESSERACT_AVAILABLE:
            try:
                pix = page.get_pixmap(dpi=300)
                img = PILImage.frombytes("RGB", [pix.width, pix.height], pix.samples)
                ocr_result = pytesseract.image_to_data(
                    img, output_type=pytesseract.Output.DICT
                )
                text = " ".join(
                    w for w in ocr_result["text"] if w.strip()
                ).strip()
                # Compute mean confidence for non-empty words
                confs = [
                    c for c, w in zip(ocr_result["conf"], ocr_result["text"])
                    if w.strip() and c != -1
                ]
                ocr_confidence = (sum(confs) / len(confs) / 100.0) if confs else 0.0
                ocr_used = True
                source = "ocr"
            except Exception as e:
                print(f"[OCR] Page {page_num + 1} OCR failed: {e}")
                source = "empty"
                text = ""
        else:
            source = "text"

        text = clean_document_text(text)
        if text:
            pages.append({
                "page": page_num + 1,
                "text": text,
                "source": source,
                "ocr_used": ocr_used,
                "ocr_confidence": ocr_confidence,
            })

    doc.close()
    return pages


def extract_docx(file_path: str) -> List[Dict]:
    """Extract text from a DOCX file. Each paragraph group maps to a virtual 'page'."""
    from docx import Document

    doc = Document(file_path)
    pages = []
    current_page_text = []
    current_page = 1
    word_count = 0
    WORDS_PER_PAGE = 400

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        words = text.split()
        current_page_text.append(text)
        word_count += len(words)

        if word_count >= WORDS_PER_PAGE:
            cleaned = clean_document_text("\n".join(current_page_text))
            if cleaned:
                pages.append({
                    "page": current_page,
                    "text": cleaned,
                    "source": "text",
                    "ocr_used": False,
                    "ocr_confidence": 1.0,
                })
            current_page += 1
            current_page_text = []
            word_count = 0

    if current_page_text:
        cleaned = clean_document_text("\n".join(current_page_text))
        if cleaned:
            pages.append({
                "page": current_page,
                "text": cleaned,
                "source": "text",
                "ocr_used": False,
                "ocr_confidence": 1.0,
            })

    return pages


def extract_pptx(file_path: str) -> List[Dict]:
    """Extract text from a PPTX file. Each slide maps to a 'page'."""
    from pptx import Presentation

    prs = Presentation(file_path)
    pages = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = " ".join(run.text for run in para.runs).strip()
                    if line:
                        texts.append(line)
        if texts:
            cleaned = clean_document_text("\n".join(texts))
            if cleaned:
                pages.append({
                    "page": slide_num,
                    "text": cleaned,
                    "source": "text",
                    "ocr_used": False,
                    "ocr_confidence": 1.0,
                })

    return pages


def extract_txt(file_path: str) -> List[Dict]:
    """Extract text from a plain TXT file. Splits into virtual pages by line count."""
    LINES_PER_PAGE = 60

    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        all_lines = [ln.rstrip() for ln in f.readlines()]

    pages = []
    current_page = 1
    for i in range(0, len(all_lines), LINES_PER_PAGE):
        chunk_lines = all_lines[i: i + LINES_PER_PAGE]
        text = clean_document_text("\n".join(chunk_lines))
        if text:
            pages.append({
                "page": current_page,
                "text": text,
                "source": "text",
                "ocr_used": False,
                "ocr_confidence": 1.0,
            })
            current_page += 1

    return pages


def extract_image(file_path: str) -> List[Dict]:
    """Extract text from an image file using OCR."""
    if not is_ocr_available():
        raise RuntimeError(
            "Tesseract OCR is not installed or available on this system. "
            "Please install Tesseract OCR (https://github.com/UB-Mannheim/tesseract/wiki) to extract text from image files."
        )
    from PIL import Image as PILImage

    img = PILImage.open(file_path)
    ocr_result = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
    raw_text = " ".join(w for w in ocr_result["text"] if w.strip()).strip()
    text = clean_document_text(raw_text)
    confs = [
        c for c, w in zip(ocr_result["conf"], ocr_result["text"])
        if w.strip() and c != -1
    ]
    ocr_confidence = (sum(confs) / len(confs) / 100.0) if confs else 0.0

    if not text:
        return []
    return [{
        "page": 1,
        "text": text,
        "source": "ocr",
        "ocr_used": True,
        "ocr_confidence": ocr_confidence,
    }]


MIME_HANDLERS = {
    "application/pdf": extract_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": extract_docx,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": extract_pptx,
    "text/plain": extract_txt,
    "image/png": extract_image,
    "image/jpeg": extract_image,
    "image/webp": extract_image,
}

EXTENSION_MAP = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".txt": "text/plain",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}


def get_mime_type(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    return EXTENSION_MAP.get(ext, "application/octet-stream")


def extract_file(file_path: str) -> List[Dict]:
    """Detect file type and extract text. Returns list of page dicts."""
    mime_type = get_mime_type(file_path)
    handler = MIME_HANDLERS.get(mime_type)
    if not handler:
        raise ValueError(f"Unsupported file type: {mime_type} ({os.path.basename(file_path)})")
    return handler(file_path)


_OCR_CHECK_CACHE = None

def is_ocr_available() -> bool:
    global _OCR_CHECK_CACHE
    if _OCR_CHECK_CACHE is not None:
        return _OCR_CHECK_CACHE
    if not TESSERACT_AVAILABLE:
        _OCR_CHECK_CACHE = False
        return False
    try:
        pytesseract.get_tesseract_version()
        _OCR_CHECK_CACHE = True
    except Exception:
        _OCR_CHECK_CACHE = False
    return _OCR_CHECK_CACHE


def compute_file_hash(file_path: str) -> str:
    """SHA-256 hash of the file contents for duplicate/version detection."""
    sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        for block in iter(lambda: f.read(65536), b""):
            sha256.update(block)
    return sha256.hexdigest()
